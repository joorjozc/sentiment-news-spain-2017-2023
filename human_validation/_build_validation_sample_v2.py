# -*- coding: utf-8 -*-
"""Build validation sample v2: representative-stratified (P1) + intensity task (P2d).

- Draws N=150 headlines from the scored 5000 pool (full 7-dim RoBERTa probability
  vectors), reflecting the real corpus emotional distribution (soft 60/40).
- Produces a HIDDEN answer key (probabilities + dominant) kept apart from annotators.
- Produces two BLIND intensity-rating templates (A1 = pareja, A2 = hermano):
  4 core-emotion Likert (0-4) + single dominant-emotion dropdown per headline.
Seed = 42 (reproducible). No model outputs are shown to annotators.
"""
import sys
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

try:
    sys.stdout.reconfigure(encoding='utf-8')
except Exception:
    pass

SEED = 42
N_TOTAL = 150
N_NEUTRAL = 90            # 60% neutral (representative-ish, soft stratification)
HERE = Path(__file__).resolve().parent
POOL = Path(r"C:/Users/jjzam/Desktop/Jorge/tesis/Metodología del estudio y resultados/Ejemplo articulo previo/balanced_sample_5000_emotions.csv")

PROB7 = ['anger', 'fear', 'joy', 'sadness', 'surprise', 'disgust', 'others']
CORE4 = ['anger', 'fear', 'joy', 'sadness']

rng = np.random.default_rng(SEED)
df = pd.read_csv(POOL)
df = df.drop_duplicates(subset='item_title').reset_index(drop=True)

# dominant emotion from the 7-dim probability vector (others -> neutral)
dom = np.array(PROB7)[df[PROB7].values.argmax(1)]
df['dominant'] = pd.Series(dom).replace({'others': 'neutral'})

# ---- representative-stratified draw ----
neutral_pool = df[df['dominant'] == 'neutral']
emo_pool = df[df['dominant'] != 'neutral']
avail = emo_pool['dominant'].value_counts()
n_emo = N_TOTAL - N_NEUTRAL  # 60 emotional

# proportional allocation across available emotional classes, floor 3 where present
alloc = {}
for emo, cnt in avail.items():
    alloc[emo] = max(3, round(n_emo * cnt / avail.sum()))
# trim/expand to exactly n_emo
while sum(alloc.values()) > n_emo:
    k = max(alloc, key=alloc.get); alloc[k] -= 1
while sum(alloc.values()) < n_emo:
    k = max(avail.to_dict(), key=lambda e: avail[e]); alloc[k] += 1
print('[alloc] neutral', N_NEUTRAL, '| emotional', alloc)

picks = [neutral_pool.sample(N_NEUTRAL, random_state=SEED)]
for emo, k in alloc.items():
    sub = emo_pool[emo_pool['dominant'] == emo]
    picks.append(sub.sample(min(k, len(sub)), random_state=SEED))
sample = pd.concat(picks).sample(frac=1, random_state=SEED).reset_index(drop=True)
sample.insert(0, 'ID', range(1, len(sample) + 1))
print('[sample] n =', len(sample), '| dominant dist:', sample['dominant'].value_counts().to_dict())

# ---- hidden answer key (NOT shown to annotators) ----
key_cols = ['ID', 'item_title', 'categoria', 'anio', 'dominant'] + PROB7
key = sample[key_cols].rename(columns={'item_title': 'Titular', 'categoria': 'categoria', 'anio': 'anio'})
key.to_csv(HERE / '_answer_key_v2.csv', index=False, encoding='utf-8')
print('[saved] _answer_key_v2.csv (hidden)')

# ---- blind intensity templates ----
LIKERT = '"0,1,2,3,4"'
EMO_DROP = '"' + ','.join(['anger', 'fear', 'joy', 'sadness', 'surprise', 'disgust', 'neutral']) + '"'
hdr_fill = PatternFill('solid', fgColor='1F4E78'); hdr_font = Font(bold=True, color='FFFFFF', size=11)
thin = Side(border_style='thin', color='AAAAAA'); border = Border(thin, thin, thin, thin)
alt = PatternFill('solid', fgColor='F2F2F2')

def build_template(tag, label):
    wb = Workbook()
    wi = wb.active; wi.title = 'Instrucciones'
    wi['A1'] = f'Validación de emociones (intensidad) — {label}'
    wi['A1'].font = Font(bold=True, size=14, color='1F4E78')
    intro = (
        "Tu tarea: leer cada titular y valorar, SOLO con lo que el titular transmite "
        "(sin buscar la noticia), DOS cosas:\n\n"
        "1) INTENSIDAD de cada una de las 4 emociones, en escala 0–4:\n"
        "     0 = nada · 1 = leve · 2 = moderada · 3 = alta · 4 = muy alta\n"
        "   Rellena las columnas Ira, Miedo, Alegría, Tristeza (un número 0–4 en cada una).\n"
        "   Un titular puede tener varias emociones a la vez, o todas a 0 si es neutral.\n\n"
        "2) EMOCIÓN DOMINANTE: elige UNA etiqueta del desplegable (las 7 opciones). "
        "Si no transmite ninguna emoción clara, elige 'neutral'.\n\n"
        "Reglas:\n"
        "  - NO dejes celdas vacías (ni las de intensidad ni la dominante).\n"
        "  - Anota tu percepción espontánea; no hay respuestas 'correctas'.\n"
        "  - NO mires resultados de modelos ni otras anotaciones.\n"
        "  - Tiempo aprox.: 45–60 min en una sola sesión.\n\n"
        "Definiciones: Ira (indignación, hostilidad, conflicto) · Miedo (amenaza, riesgo, alarma) · "
        "Alegría (celebración, éxito, buena noticia) · Tristeza (pérdida, declive, sufrimiento)."
    )
    wi['A3'] = intro
    wi['A3'].alignment = Alignment(wrap_text=True, vertical='top')
    wi.row_dimensions[3].height = 300
    wi.column_dimensions['A'].width = 110

    ws = wb.create_sheet('Titulares')
    hdrs = ['ID', 'Año', 'Categoría', 'Titular',
            'Ira (0-4)', 'Miedo (0-4)', 'Alegría (0-4)', 'Tristeza (0-4)',
            'Emoción dominante', 'Notas (opcional)']
    for c, h in enumerate(hdrs, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = hdr_font; cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = border
    ws.row_dimensions[1].height = 30
    for i, r in sample.iterrows():
        row = i + 2
        ws.cell(row=row, column=1, value=int(r['ID'])).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=2, value=int(r['anio']) if pd.notna(r['anio']) else None).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=3, value=r['categoria']).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=4, value=r['item_title']).alignment = Alignment(wrap_text=True, vertical='center')
        for c in range(5, 9):
            ws.cell(row=row, column=c).alignment = Alignment(horizontal='center')
        ws.cell(row=row, column=9).alignment = Alignment(horizontal='center')
        for c in range(1, 11):
            ws.cell(row=row, column=c).border = border
            if row % 2 == 0:
                ws.cell(row=row, column=c).fill = alt
    widths = [6, 7, 14, 80, 11, 11, 12, 12, 16, 26]
    for c, w in enumerate(widths, 1):
        ws.column_dimensions[chr(64 + c)].width = w
    dv_int = DataValidation(type='list', formula1=LIKERT, allow_blank=False)
    dv_int.prompt = 'Intensidad 0–4'
    ws.add_data_validation(dv_int); dv_int.add(f'E2:H{len(sample)+1}')
    dv_dom = DataValidation(type='list', formula1=EMO_DROP, allow_blank=False)
    dv_dom.prompt = 'Emoción dominante percibida'
    ws.add_data_validation(dv_dom); dv_dom.add(f'I2:I{len(sample)+1}')
    ws.freeze_panes = 'A2'
    out = HERE / f'Anotacion_Humana_v2_{tag}.xlsx'
    wb.save(out)
    print(f'[saved] {out.name}  (n={len(sample)}, ciega)')

build_template('A1_Millenial', 'A1 (Millenial)')
build_template('A2_Gen Z', 'A2 (Gen Z)')
print('[done]')
